"""Build the interactive Excel BI dashboard (openpyxl).

 * Dashboard sheet: route drop-down filter drives KPI formulas AND two live charts (daily trend, hourly)
 * 8 charts, KPI tiles, Executive Summary sheet, filterable Data sheet
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import PROC_DIR, EXCEL_PATH

NAVY, TEAL, LIGHT = "0B3D91", "1FB5AD", "EAF2FB"
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor=NAVY)


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")


def write_table(ws, top, left, df, title=None):
    if title:
        ws.cell(row=top - 1, column=left, value=title).font = Font(bold=True, color=NAVY)
    for j, col in enumerate(df.columns):
        ws.cell(row=top, column=left + j, value=col)
    style_header(ws, top, left, left + len(df.columns) - 1)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            ws.cell(row=top + i, column=left + j, value=v.item() if hasattr(v, "item") else v)
    return top + 1, top + len(df)


def bar(ws, anchor, title, cats, vals, horizontal=False, w=15, h=7.5, colour=NAVY):
    ch = BarChart(); ch.type = "bar" if horizontal else "col"; ch.title = title
    ch.add_data(vals, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None; ch.width, ch.height = w, h
    ch.series[0].graphicalProperties.solidFill = colour
    ws.add_chart(ch, anchor)


def main():
    df = pd.read_csv(PROC_DIR / "trips_clean.csv")
    comp = df[df.status == "Completed"]
    wb = Workbook()

    # ---------------- Data sheet ----------------
    wd = wb.active; wd.title = "Data"
    cols = ["trip_id", "date", "route_name", "bus_id", "departure_hour", "day_name", "week", "weather",
            "traffic_index", "passengers", "capacity", "occupancy_pct", "fare_pkr", "revenue_pkr",
            "delay_min", "status"]
    d = df[cols].copy(); d["delay_min"] = d["delay_min"].fillna("")
    d["date"] = pd.to_datetime(d["date"]).dt.date
    wd.append(cols)
    for r in d.itertuples(index=False):
        wd.append([v.item() if hasattr(v, "item") else v for v in r])
    style_header(wd, 1, 1, len(cols)); wd.freeze_panes = "A2"; wd.auto_filter.ref = wd.dimensions
    for i in range(1, len(cols) + 1):
        wd.column_dimensions[get_column_letter(i)].width = 16
    N = len(d) + 1
    rng = lambda col: f"Data!${get_column_letter(cols.index(col) + 1)}$2:${get_column_letter(cols.index(col) + 1)}${N}"

    # ---------------- Analysis (chart source tables) ----------------
    wa = wb.create_sheet("Analysis")
    route_rev = df.groupby("route_name", as_index=False).agg(Revenue=("revenue_pkr", "sum")).sort_values("Revenue")
    route_delay = comp.groupby("route_name", as_index=False).agg(AvgDelay=("delay_min", "mean")).round(1).sort_values("AvgDelay")
    status = df.status.value_counts().rename_axis("Status").reset_index(name="Trips")
    order = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    dow = df.groupby("day_name").passengers.sum().reindex(order).rename_axis("Day").reset_index(name="Passengers")
    weekly = df.groupby("week", as_index=False).agg(Revenue=("revenue_pkr", "sum")); weekly["week"] = "W" + weekly.week.astype(str)
    wx = comp.groupby("weather", as_index=False).agg(AvgDelay=("delay_min", "mean")).round(1)
    t = {}
    t["rev"] = write_table(wa, 2, 1, route_rev, "Revenue by route")
    t["dly"] = write_table(wa, 2, 4, route_delay, "Avg delay by route")
    t["sts"] = write_table(wa, 2, 7, status, "Trip status")
    t["dow"] = write_table(wa, 2, 10, dow, "Passengers by weekday")
    t["wk"] = write_table(wa, 2, 13, weekly, "Weekly revenue")
    t["wx"] = write_table(wa, 2, 16, wx, "Delay by weather")
    # live (formula) tables driven by dashboard filter
    dates = sorted(df.date.unique()); hours = sorted(df.departure_hour.unique())
    wa.cell(row=18, column=1, value="Daily passengers (filtered)").font = Font(bold=True, color=NAVY)
    wa.cell(row=19, column=1, value="Date"); wa.cell(row=19, column=2, value="Passengers"); style_header(wa, 19, 1, 2)
    crit = "IF(Dashboard!$C$4=\"All\",\"*\",Dashboard!$C$4)"
    for i, dt in enumerate(dates):
        r = 20 + i
        wa.cell(row=r, column=1, value=dt)
        wa.cell(row=r, column=2, value=f'=SUMIFS({rng("passengers")},{rng("date")},A{r},{rng("route_name")},{crit})')
    d_end = 19 + len(dates)
    wa.cell(row=18, column=4, value="Passengers by hour (filtered)").font = Font(bold=True, color=NAVY)
    wa.cell(row=19, column=4, value="Hour"); wa.cell(row=19, column=5, value="Passengers"); style_header(wa, 19, 4, 5)
    for i, hr in enumerate(hours):
        r = 20 + i
        wa.cell(row=r, column=4, value=int(hr))
        wa.cell(row=r, column=5, value=f'=SUMIFS({rng("passengers")},{rng("departure_hour")},D{r},{rng("route_name")},{crit})')
    h_end = 19 + len(hours)
    for i in range(1, 18):
        wa.column_dimensions[get_column_letter(i)].width = 22

    # ---------------- Dashboard ----------------
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B1:R2"); ws["B1"] = "SMART PUBLIC TRANSPORT - BI DASHBOARD"
    ws["B1"].font = Font(size=20, bold=True, color="FFFFFF"); ws["B1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws["B4"] = "Route filter:"; ws["B4"].font = Font(bold=True)
    ws["C4"] = "All"; ws["C4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.merge_cells("C4:E4")
    ws["F4"] = "<- pick a route to refresh KPIs and the two live charts"; ws["F4"].font = Font(italic=True, color="888888")
    dv = DataValidation(type="list", formula1='"All,' + ",".join(sorted(df.route_name.unique())) + '"', allow_blank=False)
    # list length > 255 chars is not allowed in inline lists -> use a range instead
    wa["T2"] = "Route list"; wa["T2"].font = Font(bold=True)
    for i, nme in enumerate(["All"] + sorted(df.route_name.unique())):
        wa.cell(row=3 + i, column=20, value=nme)
    dv = DataValidation(type="list", formula1=f"=Analysis!$T$3:$T${3 + len(df.route_name.unique())}", allow_blank=False)
    ws.add_data_validation(dv); dv.add("C4")

    f = lambda rngc, extra="": f'IF(Dashboard!$C$4="All","*",Dashboard!$C$4)'
    kpis = [
        ("Total Passengers", f'=SUMIFS({rng("passengers")},{rng("route_name")},{crit})', "#,##0"),
        ("Total Revenue (PKR)", f'=SUMIFS({rng("revenue_pkr")},{rng("route_name")},{crit})', "#,##0"),
        ("Trip Completion", f'=COUNTIFS({rng("status")},"Completed",{rng("route_name")},{crit})/COUNTIFS({rng("route_name")},{crit})', "0.0%"),
        ("Avg Delay (min)", f'=AVERAGEIFS({rng("delay_min")},{rng("status")},"Completed",{rng("route_name")},{crit})', "0.0"),
        ("Avg Occupancy", f'=AVERAGEIFS({rng("occupancy_pct")},{rng("status")},"Completed",{rng("route_name")},{crit})/100', "0.0%"),
    ]
    for i, (lab, fm, nf) in enumerate(kpis):
        c0 = 2 + i * 3
        ws.merge_cells(start_row=6, start_column=c0, end_row=6, end_column=c0 + 2)
        ws.merge_cells(start_row=7, start_column=c0, end_row=8, end_column=c0 + 2)
        a, b = ws.cell(row=6, column=c0, value=lab), ws.cell(row=7, column=c0, value=fm)
        a.font = Font(bold=True, color="FFFFFF"); a.fill = PatternFill("solid", fgColor=TEAL); a.alignment = Alignment(horizontal="center")
        b.number_format = nf; b.font = Font(size=18, bold=True, color=NAVY); b.fill = PatternFill("solid", fgColor=LIGHT)
        b.alignment = Alignment(horizontal="center", vertical="center")
    for i in range(1, 20):
        ws.column_dimensions[get_column_letter(i)].width = 9

    A = lambda k, c: (wa, t[k][0], t[k][1], c)
    def ref(k, col, hdr=False):
        s, e = t[k]; return Reference(wa, min_col=col, min_row=s - 1 if hdr else s, max_row=e)
    # 1-3 row of charts
    bar(ws, "B10", "1. Revenue by route (PKR)", ref("rev", 1), ref("rev", 2, True), horizontal=True, w=16)
    lc = LineChart(); lc.title = "2. Daily passengers (live filter)"; lc.width, lc.height = 16, 7.5
    lc.add_data(Reference(wa, min_col=2, min_row=19, max_row=d_end), titles_from_data=True)
    lc.set_categories(Reference(wa, min_col=1, min_row=20, max_row=d_end)); lc.legend = None
    ws.add_chart(lc, "K10")
    bar(ws, "B26", "3. Passengers by hour (live filter)", Reference(wa, min_col=4, min_row=20, max_row=h_end),
        Reference(wa, min_col=5, min_row=19, max_row=h_end), colour=TEAL, w=16)
    pc = PieChart(); pc.title = "4. Trip status"; pc.width, pc.height = 16, 7.5
    pc.add_data(ref("sts", 2, True), titles_from_data=True); pc.set_categories(ref("sts", 1)); ws.add_chart(pc, "K26")
    bar(ws, "B42", "5. Avg delay by route (min)", ref("dly", 4), ref("dly", 5, True), horizontal=True, w=16, colour="E07A5F")
    bar(ws, "K42", "6. Passengers by weekday", ref("dow", 10), ref("dow", 11, True), w=16)
    l2 = LineChart(); l2.title = "7. Weekly revenue (PKR)"; l2.width, l2.height = 16, 7.5
    l2.add_data(ref("wk", 14, True), titles_from_data=True); l2.set_categories(ref("wk", 13)); l2.legend = None
    ws.add_chart(l2, "B58")
    bar(ws, "K58", "8. Avg delay by weather (min)", ref("wx", 16), ref("wx", 17, True), w=16, colour="E07A5F")

    # ---------------- Executive Summary ----------------
    we = wb.create_sheet("Executive Summary", 1); we.sheet_view.showGridLines = False
    top_route = df.groupby("route_name").revenue_pkr.sum().idxmax()
    peak = int(df.groupby("departure_hour").passengers.sum().idxmax())
    worst = comp.groupby("route_name").delay_min.mean().idxmax()
    lines = [
        ("EXECUTIVE SUMMARY", True),
        (f"Network carried {df.passengers.sum():,} passengers and earned PKR {df.revenue_pkr.sum():,.0f} over 90 days.", False),
        (f"Trip completion rate: {(df.status == 'Completed').mean() * 100:.1f}%. Average delay: {comp.delay_min.mean():.1f} min per trip.", False),
        (f"Most profitable route: {top_route}. Peak travel hour: {peak}:00.", False),
        (f"Highest average delay: {worst} ({comp.groupby('route_name').delay_min.mean().max():.1f} min).", False),
        ("RECOMMENDATIONS", True),
        ("1. Add buses on the top routes during 8-9 AM and 5-6 PM peaks.", False),
        ("2. Reduce frequency on route-hours with occupancy below 25% (see SQL Q14).", False),
        ("3. Use the ML delay-risk model to pre-warn passengers on rainy peak-hour trips.", False),
        ("4. Prioritise maintenance for buses older than 8 years (higher delay and cancellation).", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=2):
        c = we.cell(row=i, column=2, value=txt); c.font = Font(bold=bold, size=14 if bold else 11, color=NAVY if bold else "000000")
    we.column_dimensions["B"].width = 120

    wb.save(EXCEL_PATH)
    print("Excel dashboard saved:", EXCEL_PATH)


if __name__ == "__main__":
    main()
